import streamlit as st
import json
import re
from urllib.parse import urlparse, parse_qs
from pypdf import PdfReader
from pptx import Presentation
from google import genai

try:
    from youtube_transcript_api import YouTubeTranscriptApi
except ImportError:
    YouTubeTranscriptApi = None


MODEL = "gemini-3.5-flash-lite"
MAX_SOURCES = 5


st.set_page_config(
    page_title="StudyMate AI",
    page_icon="📚",
    layout="wide"
)


st.markdown("""
<style>
.stApp {
    background: linear-gradient(
        135deg,
        #0b1020 0%,
        #111a30 55%,
        #17243d 100%
    );
    color: #e8edf8;
}

.block-container {
    max-width: 1200px;
    padding-top: 2.2rem;
    padding-bottom: 4rem;
}

h1, h2, h3 {
    color: #f1f5ff !important;
    letter-spacing: -0.4px;
}

.stApp p,
.stApp label,
.stMarkdown {
    color: #e2e8f0;
}

[data-testid="stCaptionContainer"] {
    color: #b8c5dc !important;
}

.stButton > button {
    background: linear-gradient(
        135deg,
        #6366f1,
        #4f46e5
    ) !important;
    color: #ffffff !important;
    border: 1px solid #818cf8 !important;
    border-radius: 10px !important;
    padding: 0.55rem 1.1rem;
    font-weight: 600;
    transition: all 0.2s ease;
}

.stButton > button:hover {
    background: #4338ca !important;
    border-color: #a5b4fc !important;
    transform: translateY(-1px);
}

[data-testid="stMetric"] {
    background: rgba(30, 41, 70, 0.9);
    border: 1px solid rgba(129, 140, 248, 0.4);
    border-radius: 14px;
    padding: 18px;
}

[data-testid="stMetricLabel"] {
    color: #cbd5e1 !important;
}

[data-testid="stMetricValue"] {
    color: #a5b4fc !important;
}

hr {
    border-color: rgba(148, 163, 184, 0.25) !important;
}

[data-testid="stFileUploader"] {
    background: rgba(30, 41, 70, 0.7);
    border: 1px dashed #818cf8;
    border-radius: 12px;
    padding: 12px;
}

.stTextInput input,
.stTextArea textarea {
    background: #f8fafc !important;
    color: #0f172a !important;
    border: 2px solid #818cf8 !important;
    border-radius: 10px !important;
}

.stTextInput input::placeholder,
.stTextArea textarea::placeholder {
    color: #475569 !important;
    opacity: 1 !important;
}

[data-testid="stChatInput"] {
    background: #f8fafc !important;
    border: 2px solid #818cf8 !important;
    border-radius: 16px !important;
    box-shadow: 0 4px 20px rgba(0, 0, 0, 0.28);
}

[data-testid="stChatInput"] textarea {
    background: #f8fafc !important;
    color: #0f172a !important;
    -webkit-text-fill-color: #0f172a !important;
    caret-color: #4f46e5 !important;
    border: none !important;
}

[data-testid="stChatInput"] textarea::placeholder {
    color: #475569 !important;
    -webkit-text-fill-color: #475569 !important;
    opacity: 1 !important;
}

[data-testid="stChatInput"]:focus-within {
    border-color: #6366f1 !important;
    box-shadow: 0 0 0 3px rgba(99, 102, 241, 0.25);
}

[data-testid="stChatInput"] button {
    color: #4f46e5 !important;
}

.stRadio {
    background: rgba(30, 41, 70, 0.7);
    border: 1px solid rgba(129, 140, 248, 0.35);
    border-radius: 12px;
    padding: 12px 16px;
}

.stRadio label {
    color: #f1f5f9 !important;
}

[data-testid="stExpander"] {
    background: rgba(30, 41, 70, 0.65);
    border: 1px solid rgba(129, 140, 248, 0.3);
    border-radius: 12px;
}

.stProgress > div > div > div {
    background: #818cf8 !important;
}

[data-testid="stChatMessage"] {
    background: rgba(30, 41, 70, 0.6);
    border: 1px solid rgba(129, 140, 248, 0.18);
    border-radius: 12px;
    padding: 12px;
}

.stButton > button,
[data-testid="stChatInput"] {
    transition: all 0.2s ease;
}
</style>
""", unsafe_allow_html=True)


@st.cache_resource
def get_client():
    return genai.Client(
        api_key=st.secrets["GEMINI_API_KEY"]
    )


def get_youtube_video_id(url):
    """Extract a YouTube video ID from common YouTube URLs."""

    try:
        parsed = urlparse(url.strip())
        host = parsed.netloc.lower().split(":")[0]

        if host in ("youtu.be", "www.youtu.be"):
            video_id = parsed.path.strip("/").split("/")[0]

        elif host in (
            "youtube.com",
            "www.youtube.com",
            "m.youtube.com",
            "music.youtube.com"
        ):
            if parsed.path == "/watch":
                video_id = parse_qs(
                    parsed.query
                ).get("v", [""])[0]

            elif parsed.path.startswith(("/embed/", "/shorts/")):
                video_id = parsed.path.strip("/").split("/")[1]

            else:
                return None

        else:
            return None

        if re.fullmatch(r"[A-Za-z0-9_-]{11}", video_id):
            return video_id

        return None

    except Exception:
        return None


def get_youtube_transcript(video_id):
    """Fetch an available transcript using youtube-transcript-api."""

    if YouTubeTranscriptApi is None:
        raise RuntimeError(
            "YouTube support is not installed. Run: "
            "python -m pip install youtube-transcript-api"
        )

    api = YouTubeTranscriptApi()

    # Current versions of youtube-transcript-api
    if hasattr(api, "fetch"):
        transcript = api.fetch(video_id)

        return "\n".join(
            getattr(item, "text", "")
            for item in transcript
        )

    # Compatibility with older versions
    if hasattr(YouTubeTranscriptApi, "get_transcript"):
        transcript = YouTubeTranscriptApi.get_transcript(
            video_id
        )

        return "\n".join(
            item.get("text", "")
            for item in transcript
        )

    raise RuntimeError(
        "Could not fetch the transcript. Please update "
        "youtube-transcript-api."
    )


st.title("📚 StudyMate AI")
st.subheader("Your Personal AI Study Companion")
st.write(
    "Upload your study material and ask questions."
)


try:
    client = get_client()

except Exception as e:
    st.error(f"API setup error: {e}")
    st.stop()


# Initialize session state

defaults = {
    "doc_name": None,
    "doc_text": "",
    "num_pages": 0,
    "messages": [],
    "quizzes_completed": 0,
    "quiz_data": [],
    "quiz_submitted": False,
    "quiz_score": 0,
    "quiz_version": 0,
    "quiz_answers": {},
    "quiz_history": [],
    "quiz_error": "",
    "source_errors": [],
    "source_signature": ""
}

for key, value in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = value


# ==========================================
# STUDY MATERIAL UPLOAD
# ==========================================

st.header("Study Material")

st.caption(
    "Upload PDFs or PowerPoint files, or add YouTube "
    "lecture links. You can use up to 5 sources."
)

uploaded_files = st.file_uploader(
    "Upload PDF or PowerPoint files",
    type=["pdf", "pptx"],
    accept_multiple_files=True,
    key="study_file_uploader"
)

youtube_links_text = st.text_area(
    "YouTube lecture links (optional)",
    placeholder=(
        "Paste one YouTube URL per line.\n"
        "Example: https://www.youtube.com/watch?v=VIDEO_ID"
    ),
    key="youtube_links_input",
    height=100
)

youtube_links = [
    line.strip()
    for line in youtube_links_text.splitlines()
    if line.strip()
]

file_count = len(uploaded_files or [])
link_count = len(youtube_links)
source_count = file_count + link_count

st.caption(
    f"Selected sources: {source_count}/{MAX_SOURCES}"
)


if source_count > MAX_SOURCES:
    st.error(
        "You can use a maximum of 5 sources in total. "
        "Remove some files or YouTube links."
    )

else:
    # Process files and links only when the selection changes.

    source_signature = " | ".join(
        [
            f"FILE:{f.name}:{f.size}"
            for f in (uploaded_files or [])
        ]
        + [
            f"YOUTUBE:{url}"
            for url in youtube_links
        ]
    )

    if source_count > 0:
        if source_signature != st.session_state.source_signature:

            all_pages = []
            total_pages = 0
            source_errors = []

            for uploaded_file in (uploaded_files or []):

                try:
                    file_extension = (
                        uploaded_file.name.lower().split(".")[-1]
                    )

                    all_pages.append(
                        f"\nSOURCE FILE: {uploaded_file.name}\n"
                    )

                    if file_extension == "pdf":
                        reader = PdfReader(uploaded_file)

                        for i, page in enumerate(
                            reader.pages,
                            start=1
                        ):
                            page_text = page.extract_text() or ""

                            all_pages.append(
                                f"{uploaded_file.name} - "
                                f"Page {i}:\n{page_text}"
                            )

                        total_pages += len(reader.pages)

                    elif file_extension == "pptx":
                        presentation = Presentation(uploaded_file)

                        for i, slide in enumerate(
                            presentation.slides,
                            start=1
                        ):
                            slide_text = "\n".join(
                                shape.text
                                for shape in slide.shapes
                                if shape.has_text_frame
                                and shape.text.strip()
                            )

                            all_pages.append(
                                f"{uploaded_file.name} - "
                                f"Slide {i}:\n{slide_text}"
                            )

                        total_pages += len(presentation.slides)

                except Exception as e:
                    source_errors.append(
                        f"Could not read {uploaded_file.name}: {e}"
                    )

            # Process YouTube links

            for url in youtube_links:

                video_id = get_youtube_video_id(url)

                if not video_id:
                    source_errors.append(
                        f"Invalid or unsupported YouTube URL: {url}"
                    )
                    continue

                try:
                    with st.spinner(
                        f"Reading YouTube transcript: {url}"
                    ):
                        transcript_text = get_youtube_transcript(
                            video_id
                        )

                    if not transcript_text.strip():
                        source_errors.append(
                            f"No readable transcript found for: {url}"
                        )
                        continue

                    all_pages.append(
                        f"\nYOUTUBE VIDEO: {url}\n"
                        f"TRANSCRIPT:\n{transcript_text}"
                    )

                except Exception as e:
                    source_errors.append(
                        f"Could not read YouTube transcript for "
                        f"{url}: {e}"
                    )

            combined_text = "\n\n".join(all_pages).strip()

            st.session_state.doc_name = (
                ", ".join(
                    [f.name for f in (uploaded_files or [])]
                    + youtube_links
                )
            )

            st.session_state.doc_text = combined_text
            st.session_state.num_pages = total_pages
            st.session_state.source_errors = source_errors
            st.session_state.source_signature = source_signature

            # Reset chat and quiz when sources change

            st.session_state.messages = []
            st.session_state.quiz_data = []
            st.session_state.quiz_submitted = False
            st.session_state.quiz_answers = {}
            st.session_state.quiz_score = 0
            st.session_state.quiz_error = ""

            st.session_state.quiz_version += 1

    else:
        # Clear old material when all sources are removed.

        if st.session_state.source_signature:
            st.session_state.doc_name = None
            st.session_state.doc_text = ""
            st.session_state.num_pages = 0
            st.session_state.messages = []
            st.session_state.quiz_data = []
            st.session_state.quiz_submitted = False
            st.session_state.quiz_answers = {}
            st.session_state.quiz_score = 0
            st.session_state.quiz_error = ""
            st.session_state.source_errors = []
            st.session_state.source_signature = ""


# Display any file or transcript errors.

for error in st.session_state.source_errors:
    st.warning(error)


document_text = st.session_state.get("doc_text", "")
has_doc = bool(document_text.strip())


if has_doc:

    st.success(
        f"Study material ready: {st.session_state.doc_name}"
    )

    st.write(
        "Total PDF pages and PowerPoint slides:",
        st.session_state.num_pages
    )

    with st.expander("View extracted study material"):
        st.text(document_text[:20000])

else:
    st.info(
        "Upload a PDF or PowerPoint, or provide a YouTube "
        "link with an available transcript."
    )


# ==========================================
# AI TUTOR
# ==========================================

st.divider()
st.header("AI Tutor")

st.caption(
    "Ask questions about your uploaded study materials."
)

question = st.chat_input(
    "Ask a question about your study material...",
    disabled=not has_doc
)


if question and has_doc:

    st.session_state.messages.append({
        "role": "user",
        "content": question
    })

    prompt = f"""
You are StudyMate AI, a helpful study tutor.

Answer the student's question using ONLY the supplied
study material. Explain concepts clearly and accurately.

If the answer is not in the material, say clearly that it
is not available in the supplied study material.

Do not invent facts or page references.
Mention the relevant page, slide, or source when possible.

STUDY MATERIAL:
{document_text[:30000]}

STUDENT QUESTION:
{question}
"""

    with st.spinner("Preparing your answer..."):

        try:
            response = client.models.generate_content(
                model=MODEL,
                contents=prompt
            )

            answer = (
                response.text
                or "No answer was returned. Please try again."
            )

        except Exception as e:
            answer = f"AI request failed: {e}"

    st.session_state.messages.append({
        "role": "assistant",
        "content": answer
    })


if st.session_state.messages:

    for message in st.session_state.messages:

        with st.chat_message(message["role"]):
            st.markdown(message["content"])


# ==========================================
# LEARNING DASHBOARD
# ==========================================

st.divider()

st.header("Learning Dashboard")

st.caption(
    "Your personal study activity and assessment overview."
)

questions_asked = sum(
    1
    for message in st.session_state.messages
    if message["role"] == "user"
)

pages_studied = st.session_state.get("num_pages", 0)

col1, col2, col3 = st.columns(3)

with col1:
    st.metric("Pages Uploaded", pages_studied)

with col2:
    st.metric("Questions Asked", questions_asked)

with col3:
    st.metric(
        "Assessments Completed",
        st.session_state.quizzes_completed
    )

st.divider()


# ==========================================
# KNOWLEDGE ASSESSMENT
# ==========================================

st.header("Knowledge Assessment")

st.caption(
    "Test your understanding of the uploaded study material."
)


if not has_doc:

    st.info(
        "Upload study material to begin your assessment."
    )

else:

    # Adapt difficulty using the previous quiz score.

    quiz_history = st.session_state.quiz_history

    if quiz_history:

        last_result = quiz_history[-1]

        last_percentage = round(
            (last_result["score"] / last_result["total"]) * 100
        )

        if last_percentage >= 80:

            difficulty = "Hard"

            difficulty_instruction = (
                "Create challenging questions that test deeper "
                "understanding, application, and analysis."
            )

        elif last_percentage >= 40:

            difficulty = "Medium"

            difficulty_instruction = (
                "Create moderately difficult questions that test "
                "conceptual understanding and application."
            )

        else:

            difficulty = "Easy"

            difficulty_instruction = (
                "Create straightforward questions that reinforce "
                "fundamental concepts."
            )

    else:

        difficulty = "Medium"

        difficulty_instruction = (
            "Create moderately difficult questions that test "
            "conceptual understanding."
        )

    st.caption(
        f"Next assessment difficulty: {difficulty}"
    )

    # Generate a new quiz.

    if st.button(
        "Generate New Quiz",
        type="primary",
        key="generate_quiz"
    ):

        quiz_prompt = f"""
Create exactly 5 multiple-choice questions using ONLY
the study material provided below.

Return ONLY a valid JSON array.
Do not use markdown or code fences.

Required format:
[
  {{
    "question": "Question text",
    "options": {{
      "A": "First option",
      "B": "Second option",
      "C": "Third option",
      "D": "Fourth option"
    }},
    "correct_answer": "A",
    "explanation": "A short explanation based on the material"
  }}
]

Rules:
1. Each question must have exactly four options: A, B, C, D.
2. Questions must test understanding of the supplied material.
3. correct_answer must be only A, B, C, or D.
4. Explanations must be short and educational.
5. Do not reveal answers outside the JSON structure.
6. Do not invent facts outside the study material.

Difficulty level: {difficulty}

Difficulty instructions:
{difficulty_instruction}

STUDY MATERIAL:
{document_text[:20000]}
"""

        with st.spinner("Preparing your assessment..."):

            try:

                response = client.models.generate_content(
                    model=MODEL,
                    contents=quiz_prompt
                )

                raw_text = (response.text or "").strip()

                # Remove Markdown code fences.

                raw_text = re.sub(
                    r"^```(?:json)?\s*|\s*```$",
                    "",
                    raw_text,
                    flags=re.IGNORECASE
                ).strip()

                # Extract JSON if extra text is returned.

                if not raw_text.startswith("["):

                    start = raw_text.find("[")
                    end = raw_text.rfind("]")

                    if start != -1 and end > start:
                        raw_text = raw_text[start:end + 1]

                quiz_data = json.loads(raw_text)

                valid_quiz = (
                    isinstance(quiz_data, list)
                    and len(quiz_data) == 5
                    and all(
                        isinstance(q, dict)
                        and isinstance(q.get("question"), str)
                        and isinstance(q.get("options"), dict)
                        and all(
                            key in q["options"]
                            and isinstance(q["options"][key], str)
                            for key in ["A", "B", "C", "D"]
                        )
                        and q.get("correct_answer") in {
                            "A", "B", "C", "D"
                        }
                        and isinstance(q.get("explanation"), str)
                        for q in quiz_data
                    )
                )

                if not valid_quiz:

                    st.session_state.quiz_error = (
                        "The AI returned an unexpected quiz format. "
                        "Please click Generate New Quiz again."
                    )

                else:

                    st.session_state.quiz_data = quiz_data
                    st.session_state.quiz_submitted = False
                    st.session_state.quiz_score = 0
                    st.session_state.quiz_answers = {}
                    st.session_state.quiz_version += 1
                    st.session_state.quiz_error = ""

            except Exception as e:

                st.session_state.quiz_error = (
                    f"Quiz generation failed: "
                    f"{type(e).__name__}: {e}"
                )


    if st.session_state.quiz_error:
        st.error(st.session_state.quiz_error)


    quiz_data = st.session_state.quiz_data


    if quiz_data:

        st.subheader("Assessment Questions")

        st.caption(
            "Select one answer for each question, then submit."
        )

        for i, q in enumerate(quiz_data):

            st.markdown(
                f"**Question {i + 1} of {len(quiz_data)}**"
            )

            st.write(q["question"])

            option_keys = ["A", "B", "C", "D"]

            selected = st.radio(
                f"Choose your answer for question {i + 1}",
                options=option_keys,
                format_func=lambda key, opts=q["options"]: (
                    f"{key}. {opts[key]}"
                ),
                key=(
                    f"assessment_"
                    f"{st.session_state.quiz_version}_{i}"
                ),
                disabled=st.session_state.quiz_submitted,
                label_visibility="collapsed"
            )

            st.session_state.quiz_answers[i] = selected

            st.divider()


        # Submit assessment answers.

        if not st.session_state.quiz_submitted:

            if st.button(
                "Submit Assessment",
                type="primary",
                key="submit_assessment"
            ):

                score = sum(
                    1
                    for i, q in enumerate(quiz_data)
                    if st.session_state.quiz_answers.get(i)
                    == q["correct_answer"]
                )

                st.session_state.quiz_score = score
                st.session_state.quiz_submitted = True

                st.session_state.quizzes_completed += 1

                st.session_state.quiz_history.append({
                    "score": score,
                    "total": len(quiz_data)
                })

                st.rerun()


        # Display results after submission.

        if st.session_state.quiz_submitted:

            score = st.session_state.quiz_score
            total = len(quiz_data)

            percentage = round(
                (score / total) * 100
            )

            st.divider()
            st.subheader("Assessment Results")

            result_col1, result_col2 = st.columns(2)

            with result_col1:
                st.metric("Your Score", f"{score}/{total}")

            with result_col2:
                st.metric("Accuracy", f"{percentage}%")

            st.progress(percentage / 100)

            if percentage >= 80:

                st.success(
                    "Excellent performance! You have a strong "
                    "understanding of this material."
                )

            elif percentage >= 60:

                st.info(
                    "Good effort! Review the explanations below "
                    "to strengthen your understanding."
                )

            else:

                st.warning(
                    "Keep practising. Review the material and "
                    "try another assessment."
                )

            st.subheader("Answer Review")

            for i, q in enumerate(quiz_data):

                user_answer = (
                    st.session_state.quiz_answers.get(i)
                )

                correct_answer = q["correct_answer"]

                st.markdown(f"**Question {i + 1}**")
                st.write(q["question"])

                if user_answer == correct_answer:

                    st.success(
                        f"Correct: {correct_answer}. "
                        f"{q['options'][correct_answer]}"
                    )

                else:

                    st.error(
                        f"Your answer: {user_answer}. "
                        f"{q['options'].get(user_answer, 'Not answered')}"
                    )

                    st.info(
                        f"Correct answer: {correct_answer}. "
                        f"{q['options'][correct_answer]}"
                    )

                st.write(
                    f"Explanation: {q['explanation']}"
                )

                st.divider()


            if st.button(
                "Start Another Assessment",
                key="another_assessment"
            ):

                st.session_state.quiz_data = []
                st.session_state.quiz_submitted = False
                st.session_state.quiz_answers = {}
                st.session_state.quiz_score = 0
                st.session_state.quiz_error = ""

                st.rerun()


# ==========================================
# LEARNING PROGRESS
# ==========================================

st.divider()

st.header("Learning Progress")

st.caption(
    "An overview of your activity in this session."
)

completed = st.session_state.quizzes_completed
history = st.session_state.quiz_history


if history:

    total_correct = sum(
        item["score"]
        for item in history
    )

    total_questions = sum(
        item["total"]
        for item in history
    )

    overall_accuracy = (
        round((total_correct / total_questions) * 100)
        if total_questions
        else 0
    )

else:

    overall_accuracy = 0


progress_col1, progress_col2 = st.columns(2)

with progress_col1:
    st.metric("Completed Assessments", completed)

with progress_col2:
    st.metric(
        "Overall Quiz Accuracy",
        f"{overall_accuracy}%"
    )


if completed > 0:

    st.progress(
        min(completed / 5, 1.0)
    )

    st.caption(
        f"{completed} assessment(s) completed. "
        "Keep practising to strengthen your knowledge."
    )

else:

    st.info(
        "Your learning statistics will update after you "
        "complete your first assessment."
    )